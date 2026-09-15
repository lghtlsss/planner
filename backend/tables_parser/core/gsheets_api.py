"""
Получение и разбор расписания напрямую через Google Sheets API v4 — без
скачивания xlsx-файла целиком.

Это ОТДЕЛЬНАЯ, самостоятельная версия получения данных для ссылок на
Google Таблицы. `core/parser.py` (openpyxl-based, работает с локальными
файлами) НЕ изменён и продолжает работать как раньше — этот модуль его
не трогает, а параллельно реализует ту же самую логику извлечения
структуры (см. docstring core/parser.py), но поверх JSON от API вместо
объекта openpyxl.Worksheet.

Почему API, а не export?format=xlsx (core/source.py:GoogleSheetsSource):
- xlsx-экспорт Google генерирует "на лету", рендеря лист со всеми
  стилями/форматированием, а затем отдаёт готовый файл с отдельного
  CDN-хоста (googleusercontent.com). На практике это может быть заметно
  медленнее обычного API-запроса и (как показала практика) иногда
  обрывается таймаутом при скачивании самого сгенерированного файла.
- Sheets API отдаёт готовый JSON с содержимым ячеек и merge-диапазонами
  напрямую, без промежуточного рендеринга в бинарный формат. Параметр
  `fields` позволяет запросить только то, что реально нужно (значения
  ячеек + merges), исключив стили/цвета/шрифты/примечания — заметно
  меньше данных и быстрее ответ.
- Параметр `ranges` позволяет запросить только нужные листы (курсы),
  не трогая, например, тяжёлый раскрашенный "Календарь" целиком — для
  него хватает маленького диапазона A1:C5, где ищется учебный год.
- Для публичных таблиц ("доступ по ссылке") достаточно простого
  API-ключа, без OAuth и без сервисного аккаунта.

Как получить API-ключ (один раз, бесплатно):
1. https://console.cloud.google.com/ -> создать проект (или выбрать
   существующий).
2. "APIs & Services" -> "Library" -> найти "Google Sheets API" -> Enable.
3. "APIs & Services" -> "Credentials" -> "Create Credentials" -> "API key".
4. (Желательно, не обязательно) "Restrict key" -> API restrictions ->
   оставить только "Google Sheets API", чтобы ключ нельзя было
   использовать для других сервисов Google.

Ключ передаётся через переменную окружения GOOGLE_SHEETS_API_KEY или
флаг командной строки --api-key (см. main.py).
"""

from __future__ import annotations

import re
from typing import Optional

import requests
from openpyxl.utils import get_column_letter

from .parser import (
    DAY_ABBREVIATIONS,
    TIME_RE,
    TITLE_RE,
    RawBuildingEntry,
    RawCell,
    RawGroupBlock,
    RawScheduleMeta,
    RawSheet,
)

SHEETS_API_BASE = "https://sheets.googleapis.com/v4/spreadsheets"
API_TIMEOUT = 30

# Просим у API только то, что реально нужно: значения ячеек (в готовом
# отображаемом виде — formattedValue) и merge-диапазоны. Без этого
# ограничения API отдаёт ещё и все стили/форматирование каждой ячейки,
# что для этой задачи бесполезно и заметно раздувает ответ.
GRID_FIELDS = "sheets(properties(title),merges,data(rowData(values(formattedValue))))"
YEAR_RE = re.compile(r"(\d{4})\s*-\s*\d{4}\s*учебный год", re.IGNORECASE)


class GoogleSheetsAPIError(Exception):
    """Ошибка обращения к Google Sheets API (плохой ключ, нет доступа,
    таблица не найдена и т. п.) — сообщение уже человекочитаемое."""


def _api_get(spreadsheet_id: str, api_key: str, params: dict) -> dict:
    url = f"{SHEETS_API_BASE}/{spreadsheet_id}"
    try:
        resp = requests.get(url, params={**params, "key": api_key}, timeout=API_TIMEOUT)
    except requests.RequestException as exc:
        raise GoogleSheetsAPIError(f"Не удалось обратиться к Google Sheets API: {exc}") from exc

    if resp.status_code != 200:
        message = resp.text
        try:
            message = resp.json()["error"]["message"]
        except Exception:  # noqa: BLE001 — тело ответа может быть каким угодно
            pass
        hint = ""
        if resp.status_code in (400, 401):
            hint = " Проверьте, что API-ключ верный и включён Google Sheets API."
        elif resp.status_code == 403:
            hint = " Похоже, таблица не открыта 'для всех, у кого есть ссылка'."
        elif resp.status_code == 404:
            hint = " Проверьте ID таблицы в ссылке — возможно, таблица удалена или ID неверный."
        raise GoogleSheetsAPIError(
            f"Google Sheets API вернул ошибку {resp.status_code}: {message}.{hint}"
        )
    return resp.json()


def list_sheet_titles(spreadsheet_id: str, api_key: str) -> list[str]:
    """Дешёвый запрос без данных ячеек — только чтобы узнать названия
    вкладок (и не тащить лишнего в основном запросе)."""
    data = _api_get(spreadsheet_id, api_key, {"fields": "sheets.properties.title"})
    return [s["properties"]["title"] for s in data.get("sheets", [])]


def fetch_grid_data(spreadsheet_id: str, api_key: str, ranges: list[str]) -> dict:
    return _api_get(
        spreadsheet_id,
        api_key,
        {"includeGridData": "true", "ranges": ranges, "fields": GRID_FIELDS},
    )


# ---------------------------------------------------------------------------
# Обёртка над JSON одного листа — даёт тот же интерфейс (значение по
# 1-индексированным (row, col), список merge-диапазонов), какой нужен
# для повторения логики извлечения структуры из core/parser.py, но не
# требует объекта openpyxl.
# ---------------------------------------------------------------------------


class ApiSheetGrid:
    def __init__(self, sheet_json: dict):
        self.title = sheet_json["properties"]["title"]
        row_data = (sheet_json.get("data") or [{}])[0].get("rowData") or []
        self._rows: list[list[Optional[str]]] = [
            [v.get("formattedValue") for v in (row.get("values") or [])] for row in row_data
        ]
        self.max_row = len(self._rows)
        self.max_col = max((len(r) for r in self._rows), default=0)

        # API отдаёт 0-индексированные полуоткрытые диапазоны; приводим
        # к 1-индексированным включительным — как в openpyxl.
        self.merges: list[tuple[int, int, int, int]] = []
        for m in sheet_json.get("merges", []):
            self.merges.append(
                (
                    m["startRowIndex"] + 1,
                    m["endRowIndex"],
                    m["startColumnIndex"] + 1,
                    m["endColumnIndex"],
                )
            )

    def value(self, row: int, col: int) -> Optional[str]:
        r_idx, c_idx = row - 1, col - 1
        if r_idx < 0 or r_idx >= len(self._rows):
            return None
        cells = self._rows[r_idx]
        if c_idx < 0 or c_idx >= len(cells):
            return None
        return cells[c_idx]


class ApiMergedIndex:
    def __init__(self, grid: ApiSheetGrid):
        self._map: dict[tuple[int, int], tuple[int, int]] = {}
        for min_row, max_row, min_col, max_col in grid.merges:
            top_left = (min_row, min_col)
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    self._map[(r, c)] = top_left

    def top_left_of(self, row: int, col: int) -> tuple[int, int]:
        return self._map.get((row, col), (row, col))

    def is_top_left(self, row: int, col: int) -> bool:
        return self.top_left_of(row, col) == (row, col)


def _effective_value(grid: ApiSheetGrid, merged: ApiMergedIndex, row: int, col: int):
    tr, tc = merged.top_left_of(row, col)
    return grid.value(tr, tc)


# ---------------------------------------------------------------------------
# Извлечение структуры — зеркалит core/parser.py по смыслу (те же опорные
# метки: "День", "Ауд.", сокращения дней недели), но читает ApiSheetGrid
# вместо openpyxl.Worksheet.
# ---------------------------------------------------------------------------


def _find_header_anchor(grid: ApiSheetGrid) -> Optional[int]:
    for r in range(1, grid.max_row + 1):
        v = grid.value(r, 2)
        if isinstance(v, str) and v.strip() == "День":
            return r
    return None


def _find_group_blocks(
    grid: ApiSheetGrid, merged: ApiMergedIndex, header_row: int, first_data_row: int
) -> list[RawGroupBlock]:
    blocks: list[RawGroupBlock] = []
    for row_idx in range(header_row, first_data_row):
        for col_idx in range(1, grid.max_col + 1):
            val = grid.value(row_idx, col_idx)
            if isinstance(val, str) and val.strip() == "Ауд.":
                room_col = col_idx
                subject_col = col_idx - 1
                building_col = col_idx + 1
                group_name_val = grid.value(row_idx, subject_col)
                if not isinstance(group_name_val, str) or not group_name_val.strip():
                    continue
                stream_label = ""
                for r in range(row_idx - 1, header_row - 1, -1):
                    v = _effective_value(grid, merged, r, subject_col)
                    if isinstance(v, str) and v.strip():
                        stream_label = v.strip()
                        break
                blocks.append(
                    RawGroupBlock(
                        group_name=group_name_val.strip(),
                        stream_label=stream_label,
                        subject_col=subject_col,
                        room_col=room_col,
                        building_col=building_col,
                    )
                )
    return blocks


def _find_day_rows(grid: ApiSheetGrid, first_data_row: int) -> list[tuple[str, int, int]]:
    result = []
    for min_row, max_row, min_col, max_col in sorted(grid.merges, key=lambda m: m[0]):
        if min_col != 2:
            continue
        val = grid.value(min_row, 2)
        if isinstance(val, str) and val.strip().lower() in DAY_ABBREVIATIONS and min_row >= first_data_row:
            result.append((val.strip(), min_row, max_row))
    return result


def _parse_metadata(grid: ApiSheetGrid, faculty: str, program: str) -> RawScheduleMeta:
    title = None
    for c in range(1, grid.max_col + 1):
        v = grid.value(1, c)
        if isinstance(v, str) and "БАКАЛАВРИАТ" in v:
            title = v
            break
    course = module = None
    p_start = p_end = None
    if title:
        m = TITLE_RE.search(title)
        if m:
            course = int(m.group(1))
            module = int(m.group(2))
            p_start, p_end = m.group(3), m.group(4)
    return RawScheduleMeta(
        faculty=faculty, program=program, course=course, module=module,
        period_start_raw=p_start, period_end_raw=p_end,
    )


def _parse_buildings(grid: ApiSheetGrid, max_row: int = 9) -> list[RawBuildingEntry]:
    seen: dict[str, str] = {}
    for r in range(1, min(max_row, grid.max_row) + 1):
        for c in range(1, grid.max_col + 1):
            v = grid.value(r, c)
            if isinstance(v, str) and v.strip().lower().startswith("корпус"):
                code = grid.value(r, c + 1)
                if isinstance(code, str) and code.strip():
                    seen[code.strip()] = v.strip()
    return [RawBuildingEntry(code=c, address=a) for c, a in seen.items()]


def extract_raw_sheet_from_api(sheet_json: dict, faculty: str, program: str) -> RawSheet:
    grid = ApiSheetGrid(sheet_json)
    merged = ApiMergedIndex(grid)

    meta = _parse_metadata(grid, faculty, program)
    buildings = _parse_buildings(grid)
    raw = RawSheet(sheet_name=grid.title, meta=meta, buildings=buildings)

    header_row = _find_header_anchor(grid)
    if header_row is None:
        raw.structural_errors.append(
            "Не найдена опорная ячейка 'День' в столбце B — структура листа не распознана."
        )
        return raw

    first_data_row = None
    for r in range(header_row + 1, grid.max_row + 1):
        v = grid.value(r, 3)
        if isinstance(v, str) and TIME_RE.match(v.strip()):
            first_data_row = r
            break
    if first_data_row is None:
        raw.structural_errors.append("Не найдена первая строка с временным интервалом в столбце C.")
        return raw

    blocks = _find_group_blocks(grid, merged, header_row, first_data_row)
    if not blocks:
        raw.structural_errors.append("Не найден ни один блок группы (по заголовку 'Ауд.').")
        return raw

    day_rows = _find_day_rows(grid, first_data_row)
    if not day_rows:
        raw.structural_errors.append("Не найден ни один день недели в столбце B.")
        return raw

    last_day_end_row = max(end for _, _, end in day_rows)

    for day_raw, row_start, row_end in day_rows:
        for r in range(row_start, row_end + 1):
            time_val = grid.value(r, 3)
            if not (isinstance(time_val, str) and TIME_RE.match(time_val.strip())):
                continue
            start_time, end_time = [t.strip() for t in time_val.split("-", 1)]
            for block in blocks:
                subject_text = _effective_value(grid, merged, r, block.subject_col)
                if subject_text is None:
                    continue
                room_text = _effective_value(grid, merged, r, block.room_col)
                building_text = _effective_value(grid, merged, r, block.building_col)
                coord = f"{get_column_letter(block.subject_col)}{r}"
                raw.cells.append(
                    RawCell(
                        group_name=block.group_name,
                        stream_label=block.stream_label,
                        day_raw=day_raw,
                        start_time=start_time,
                        end_time=end_time,
                        subject_text=subject_text,
                        room_text=room_text,
                        building_text=building_text,
                        coordinate=coord,
                    )
                )

    seen_notes: set[str] = set()
    for r in range(last_day_end_row + 1, grid.max_row + 1):
        for c in range(1, grid.max_col + 1):
            if not merged.is_top_left(r, c):
                continue
            v = grid.value(r, c)
            if isinstance(v, str) and v.strip() and v.strip() not in seen_notes:
                seen_notes.add(v.strip())
                raw.sheet_notes.append(v.strip())

    return raw


# ---------------------------------------------------------------------------
# Оркестрация: по ID таблицы -> список RawSheet по листам-курсам +
# учебный год (из листа "Календарь", без загрузки его целиком).
# ---------------------------------------------------------------------------


def build_raw_sheets_from_api(
    spreadsheet_id: str,
    api_key: str,
    faculty: str,
    program: str,
    sheet_filter_regex: str = "курс",
) -> tuple[list[RawSheet], Optional[int]]:
    all_titles = list_sheet_titles(spreadsheet_id, api_key)
    course_titles = [t for t in all_titles if re.search(sheet_filter_regex, t, re.IGNORECASE)]
    if not course_titles:
        raise GoogleSheetsAPIError(
            f"Ни один лист не подошёл под фильтр '{sheet_filter_regex}'. "
            f"Доступные листы: {', '.join(all_titles)}"
        )

    calendar_title = next((t for t in all_titles if "календар" in t.lower()), None)
    ranges = list(course_titles)
    if calendar_title:
        # Не тащим весь раскрашенный годовой календарь целиком — там
        # достаточно маленького диапазона, где написан учебный год.
        ranges.append(f"{calendar_title}!A1:C5")

    data = fetch_grid_data(spreadsheet_id, api_key, ranges)

    academic_year_start: Optional[int] = None
    raw_sheets: list[RawSheet] = []

    for sheet_json in data.get("sheets", []):
        title = sheet_json["properties"]["title"]
        if calendar_title and title == calendar_title:
            grid = ApiSheetGrid(sheet_json)
            for r in range(1, grid.max_row + 1):
                for c in range(1, grid.max_col + 1):
                    v = grid.value(r, c)
                    if isinstance(v, str):
                        m = YEAR_RE.search(v)
                        if m:
                            academic_year_start = int(m.group(1))
            continue
        raw_sheets.append(extract_raw_sheet_from_api(sheet_json, faculty, program))

    return raw_sheets, academic_year_start
