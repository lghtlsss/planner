"""
Получение и разбор расписания напрямую через Google Sheets API v4 — без
скачивания xlsx-файла. Основной путь для ссылок на Google Таблицы.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from typing import Optional

import requests
from openpyxl.utils import get_column_letter

from .parser import (
    RawBuildingEntry,
    RawCell,
    RawGroupBlock,
    RawScheduleMeta,
    RawSheet,
)

SHEETS_API_BASE = "https://sheets.googleapis.com/v4/spreadsheets"
API_TIMEOUT = 30

# Только значения ячеек и merge-диапазоны — без стилей, цветов и т. п.
GRID_FIELDS = "sheets(properties(title),merges,data(rowData(values(formattedValue))))"
YEAR_RE = re.compile(r"(\d{4})\s*-\s*\d{4}\s*учебный год", re.IGNORECASE)

_DAY_MAP = {
    "пн": "Пн", "понедельник": "Пн",
    "вт": "Вт", "вторник": "Вт",
    "ср": "Ср", "среда": "Ср",
    "чт": "Чт", "четверг": "Чт",
    "пт": "Пт", "пятница": "Пт",
    "сб": "Сб", "суббота": "Сб",
    "вс": "Вс", "воскресенье": "Вс",
}
_TIME_RE = re.compile(r"^\s*(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})\s*$")

_DEGREE_RE = re.compile(r"(БАКАЛАВРИАТ|МАГИСТРАТУРА|СПЕЦИАЛИТЕТ|АСПИРАНТУРА)", re.IGNORECASE)
_DEGREE_SHORT = {"БАКАЛАВРИАТ": "БАК", "МАГИСТРАТУРА": "МАГ", "СПЕЦИАЛИТЕТ": "СПЕЦ", "АСПИРАНТУРА": "АСП"}
_COURSE_RE = re.compile(r"(\d+)\s*курс", re.IGNORECASE)
_MODULE_RE = re.compile(r"(\d+)\s*модул", re.IGNORECASE)
_PERIOD_RE = re.compile(r"\(\s*([\d.]+)\s*[-–]\s*([\d.]+)\s*\)")
_SS_TITLE_RE = re.compile(
    r"Ф[-‑–]?т\.?\s+(?P<fac>.+?)\s+[-–—]\s+(?P<deg>БАК|МАГ|СПЕЦ|АСП)[А-ЯЁа-яё]*\.?\s*(?P<prog>.*)$",
    re.IGNORECASE,
)


class GoogleSheetsAPIError(Exception):
    """Ошибка обращения к Google Sheets API."""


@dataclass
class ApiRawSheet(RawSheet):
    degree: Optional[str] = None  # БАК / МАГ / СПЕЦ / АСП


@dataclass
class ApiFetchResult:
    raw_sheets: list
    academic_year_start: Optional[int]
    spreadsheet_title: Optional[str]
    faculty: Optional[str]
    program: Optional[str]
    content_hash: Optional[str] = None


@dataclass
class ApiPayload:
    """Скачанные из API данные + хеш содержимого (для проверки «расписание изменилось?»)."""

    spreadsheet_id: str
    spreadsheet_title: Optional[str]
    calendar_title: Optional[str]
    data: dict
    content_hash: str


def _api_get(spreadsheet_id: str, api_key: str, params: dict) -> dict:
    url = f"{SHEETS_API_BASE}/{spreadsheet_id}"
    try:
        resp = requests.get(url, params={**params, "key": api_key}, timeout=API_TIMEOUT)
    except requests.RequestException as exc:
        raise GoogleSheetsAPIError(f"Не удалось обратиться к Google Sheets API: {exc}") from exc

    if resp.status_code != 200:
        message = resp.text
        try:
            message = resp.json()["error"]["message"]
        except Exception:  # noqa: BLE001
            pass
        hint = ""
        if resp.status_code in (400, 401):
            hint = " Проверьте, что API-ключ верный и включён Google Sheets API."
        elif resp.status_code == 403:
            hint = " Похоже, таблица не открыта 'для всех, у кого есть ссылка'."
        elif resp.status_code == 404:
            hint = " Проверьте ID таблицы в ссылке — возможно, таблица удалена или ID неверный."
        raise GoogleSheetsAPIError(f"Google Sheets API вернул ошибку {resp.status_code}: {message}.{hint}")
    return resp.json()


def get_spreadsheet_overview(spreadsheet_id: str, api_key: str) -> tuple[Optional[str], list[str]]:
    """Дешёвый запрос: название таблицы и названия вкладок."""
    data = _api_get(spreadsheet_id, api_key, {"fields": "properties.title,sheets.properties.title"})
    title = (data.get("properties") or {}).get("title")
    titles = [s["properties"]["title"] for s in data.get("sheets", [])]
    return title, titles


def list_sheet_titles(spreadsheet_id: str, api_key: str) -> list[str]:
    return get_spreadsheet_overview(spreadsheet_id, api_key)[1]


def fetch_grid_data(spreadsheet_id: str, api_key: str, ranges: list[str]) -> dict:
    return _api_get(
        spreadsheet_id, api_key,
        {"includeGridData": "true", "ranges": ranges, "fields": GRID_FIELDS},
    )


def parse_spreadsheet_title(title: Optional[str]) -> tuple[Optional[str], Optional[str]]:
    """'1 модуль - Ф-т ПРАВА - БАК (Юриспруденция)' -> ('ПРАВА', 'БАК "Юриспруденция"')."""
    if not title:
        return None, None
    m = _SS_TITLE_RE.search(title)
    if not m:
        return None, None
    faculty = m.group("fac").strip()
    deg = m.group("deg").upper()
    prog = m.group("prog").strip().strip('"«»“”()[] ').strip()
    return faculty, (f'{deg} "{prog}"' if prog else deg)


# ---------------------------------------------------------------------------
# Обёртка над JSON листа (1-индексированные row/col, как в openpyxl)
# ---------------------------------------------------------------------------


class ApiSheetGrid:
    def __init__(self, sheet_json: dict):
        self.title = sheet_json["properties"]["title"]
        row_data = (sheet_json.get("data") or [{}])[0].get("rowData") or []
        self._rows: list[list[Optional[str]]] = [
            [v.get("formattedValue") for v in (row.get("values") or [])] for row in row_data
        ]
        self.max_row = len(self._rows)
        self.max_col = max((len(r) for r in self._rows), default=0)
        # API: 0-индексированные полуоткрытые диапазоны -> 1-индексированные включительные
        self.merges: list[tuple[int, int, int, int]] = [
            (m["startRowIndex"] + 1, m["endRowIndex"], m["startColumnIndex"] + 1, m["endColumnIndex"])
            for m in sheet_json.get("merges", [])
        ]

    def value(self, row: int, col: int) -> Optional[str]:
        r_idx, c_idx = row - 1, col - 1
        if r_idx < 0 or r_idx >= len(self._rows):
            return None
        cells = self._rows[r_idx]
        if c_idx < 0 or c_idx >= len(cells):
            return None
        return cells[c_idx]


class ApiMergedIndex:
    def __init__(self, grid: ApiSheetGrid):
        self._map: dict[tuple[int, int], tuple[int, int]] = {}
        for min_row, max_row, min_col, max_col in grid.merges:
            top_left = (min_row, min_col)
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    self._map[(r, c)] = top_left

    def top_left_of(self, row: int, col: int) -> tuple[int, int]:
        return self._map.get((row, col), (row, col))

    def is_top_left(self, row: int, col: int) -> bool:
        return self.top_left_of(row, col) == (row, col)


def _effective_value(grid: ApiSheetGrid, merged: ApiMergedIndex, row: int, col: int):
    tr, tc = merged.top_left_of(row, col)
    return grid.value(tr, tc)


# ---------------------------------------------------------------------------
# Извлечение структуры
# ---------------------------------------------------------------------------


def _normalize_day(v) -> Optional[str]:
    if not isinstance(v, str):
        return None
    return _DAY_MAP.get(v.strip().lower().rstrip("."))


def _parse_time_range(v) -> Optional[tuple[str, str]]:
    if not isinstance(v, str):
        return None
    m = _TIME_RE.match(v)
    if not m:
        return None
    h1, m1, h2, m2 = m.groups()
    return f"{int(h1):02d}:{m1}", f"{int(h2):02d}:{m2}"


def _find_anchor(grid: ApiSheetGrid) -> Optional[tuple[int, int, int]]:
    """(строка шапки, столбец дня, столбец времени) по заголовку 'День'."""
    for r in range(1, min(grid.max_row, 40) + 1):
        for c in range(1, grid.max_col + 1):
            v = grid.value(r, c)
            if isinstance(v, str) and v.strip().lower() == "день":
                time_col = c + 1
                for c2 in range(1, grid.max_col + 1):
                    v2 = grid.value(r, c2)
                    if isinstance(v2, str) and v2.strip().lower() == "время":
                        time_col = c2
                        break
                return r, c, time_col
    return None


def _find_group_blocks(
    grid: ApiSheetGrid, merged: ApiMergedIndex, header_row: int, first_data_row: int
) -> list[RawGroupBlock]:
    blocks: list[RawGroupBlock] = []
    for row_idx in range(header_row, first_data_row):
        for col_idx in range(1, grid.max_col + 1):
            val = grid.value(row_idx, col_idx)
            if not (isinstance(val, str) and val.strip().lower().startswith("ауд")):
                continue
            subject_col = col_idx - 1
            if subject_col < 1:
                continue
            group_name_val = grid.value(row_idx, subject_col)
            if not isinstance(group_name_val, str) or not group_name_val.strip():
                continue
            stream_label = ""
            for r in range(row_idx - 1, header_row - 1, -1):
                v = _effective_value(grid, merged, r, subject_col)
                if isinstance(v, str) and v.strip():
                    stream_label = v.strip()
                    break
            blocks.append(
                RawGroupBlock(
                    group_name=group_name_val.strip(),
                    stream_label=stream_label,
                    subject_col=subject_col,
                    room_col=col_idx,
                    building_col=col_idx + 1,
                )
            )
    return blocks


def _parse_metadata(
    grid: ApiSheetGrid, faculty: Optional[str], program: Optional[str]
) -> tuple[RawScheduleMeta, Optional[str]]:
    title = None
    degree = None
    for r in range(1, min(grid.max_row, 3) + 1):
        for c in range(1, grid.max_col + 1):
            v = grid.value(r, c)
            if not isinstance(v, str):
                continue
            dm = _DEGREE_RE.search(v)
            if dm and (_COURSE_RE.search(v) or _MODULE_RE.search(v)):
                title, degree = v, _DEGREE_SHORT[dm.group(1).upper()]
                break
        if title:
            break
    if title is None:
        for r in range(1, min(grid.max_row, 3) + 1):
            for c in range(1, grid.max_col + 1):
                v = grid.value(r, c)
                if isinstance(v, str) and _COURSE_RE.search(v) and _MODULE_RE.search(v):
                    title = v
                    break
            if title:
                break

    course = module = None
    p_start = p_end = None
    if title:
        cm, mm, pm = _COURSE_RE.search(title), _MODULE_RE.search(title), _PERIOD_RE.search(title)
        course = int(cm.group(1)) if cm else None
        module = int(mm.group(1)) if mm else None
        if pm:
            p_start, p_end = pm.group(1), pm.group(2)
    meta = RawScheduleMeta(
        faculty=faculty, program=program or degree, course=course, module=module,
        period_start_raw=p_start, period_end_raw=p_end,
    )
    return meta, degree


def _parse_buildings(grid: ApiSheetGrid, max_row: int = 9) -> list[RawBuildingEntry]:
    seen: dict[str, str] = {}
    for r in range(1, min(max_row, grid.max_row) + 1):
        for c in range(1, grid.max_col + 1):
            v = grid.value(r, c)
            if isinstance(v, str) and v.strip().lower().startswith("корпус"):
                code = grid.value(r, c + 1)
                if isinstance(code, str) and code.strip():
                    seen[code.strip()] = v.strip()
    return [RawBuildingEntry(code=c, address=a) for c, a in seen.items()]


def extract_raw_sheet_from_api(
    sheet_json: dict, faculty: Optional[str] = None, program: Optional[str] = None
) -> ApiRawSheet:
    grid = ApiSheetGrid(sheet_json)
    merged = ApiMergedIndex(grid)

    meta, degree = _parse_metadata(grid, faculty, program)
    raw = ApiRawSheet(sheet_name=grid.title, meta=meta, buildings=_parse_buildings(grid), degree=degree)

    anchor = _find_anchor(grid)
    if anchor is None:
        raw.structural_errors.append("Не найдена опорная ячейка 'День' в шапке — структура листа не распознана.")
        return raw
    header_row, day_col, time_col = anchor

    first_data_row = None
    for r in range(header_row + 1, grid.max_row + 1):
        if _parse_time_range(_effective_value(grid, merged, r, time_col)):
            first_data_row = r
            break
    if first_data_row is None:
        raw.structural_errors.append("Не найдена первая строка с временным интервалом.")
        return raw

    blocks = _find_group_blocks(grid, merged, header_row, first_data_row)
    if not blocks:
        raw.structural_errors.append("Не найден ни один блок группы (по заголовку 'Ауд.').")
        return raw

    if not any(_normalize_day(grid.value(r, day_col)) for r in range(first_data_row, grid.max_row + 1)):
        raw.structural_errors.append("Не найден ни один день недели в столбце 'День'.")
        return raw

    current_day: Optional[str] = None
    last_time_row = first_data_row
    for r in range(first_data_row, grid.max_row + 1):
        d = _normalize_day(grid.value(r, day_col))  # именно значение, не merge: метка только в верхней ячейке
        if d:
            current_day = d
        tr = _parse_time_range(_effective_value(grid, merged, r, time_col))
        if not tr or current_day is None:
            continue
        last_time_row = r
        start_time, end_time = tr
        for block in blocks:
            subject_text = _effective_value(grid, merged, r, block.subject_col)
            if subject_text is None or not str(subject_text).strip():
                continue
            raw.cells.append(
                RawCell(
                    group_name=block.group_name,
                    stream_label=block.stream_label,
                    day_raw=current_day,
                    start_time=start_time,
                    end_time=end_time,
                    subject_text=subject_text,
                    room_text=_effective_value(grid, merged, r, block.room_col),
                    building_text=_effective_value(grid, merged, r, block.building_col),
                    coordinate=f"{get_column_letter(block.subject_col)}{r}",
                )
            )

    # Примечания уровня листа — ниже последней строки со временем
    seen_notes: set[str] = set()
    for r in range(last_time_row + 1, grid.max_row + 1):
        for c in range(1, grid.max_col + 1):
            if not merged.is_top_left(r, c):
                continue
            v = grid.value(r, c)
            if isinstance(v, str) and v.strip() and v.strip() not in seen_notes:
                seen_notes.add(v.strip())
                raw.sheet_notes.append(v.strip())
    return raw


# ---------------------------------------------------------------------------
# Оркестрация
# ---------------------------------------------------------------------------


def fetch_payload(
    spreadsheet_id: str, api_key: str, sheet_filter_regex: str = "курс"
) -> ApiPayload:
    """Только скачивание (2 запроса к API) и хеш — без разбора."""
    ss_title, all_titles = get_spreadsheet_overview(spreadsheet_id, api_key)

    course_titles = [t for t in all_titles if re.search(sheet_filter_regex, t, re.IGNORECASE)]
    if not course_titles:
        raise GoogleSheetsAPIError(
            f"Ни один лист не подошёл под фильтр '{sheet_filter_regex}'. "
            f"Доступные листы: {', '.join(all_titles)}"
        )

    calendar_title = next((t for t in all_titles if "календар" in t.lower()), None)
    ranges = list(course_titles)
    if calendar_title:
        ranges.append(f"{calendar_title}!A1:H5")  # только где написан учебный год

    data = fetch_grid_data(spreadsheet_id, api_key, ranges)
    canonical = json.dumps(data, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    content_hash = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    return ApiPayload(spreadsheet_id, ss_title, calendar_title, data, content_hash)


def parse_payload(
    payload: ApiPayload, faculty: Optional[str] = None, program: Optional[str] = None
) -> ApiFetchResult:
    """Разбор уже скачанных данных (без обращений к сети)."""
    derived_faculty, derived_program = parse_spreadsheet_title(payload.spreadsheet_title)
    faculty = faculty or derived_faculty
    program = program or derived_program

    academic_year_start: Optional[int] = None
    raw_sheets: list[RawSheet] = []
    for sheet_json in payload.data.get("sheets", []):
        title = sheet_json["properties"]["title"]
        if payload.calendar_title and title == payload.calendar_title:
            grid = ApiSheetGrid(sheet_json)
            for r in range(1, grid.max_row + 1):
                for c in range(1, grid.max_col + 1):
                    v = grid.value(r, c)
                    if isinstance(v, str):
                        m = YEAR_RE.search(v)
                        if m:
                            academic_year_start = int(m.group(1))
            continue
        raw_sheets.append(extract_raw_sheet_from_api(sheet_json, faculty, program))

    return ApiFetchResult(
        raw_sheets, academic_year_start, payload.spreadsheet_title, faculty, program,
        payload.content_hash,
    )


def fetch_schedule_from_api(
    spreadsheet_id: str,
    api_key: str,
    faculty: Optional[str] = None,
    program: Optional[str] = None,
    sheet_filter_regex: str = "курс",
) -> ApiFetchResult:
    payload = fetch_payload(spreadsheet_id, api_key, sheet_filter_regex)
    return parse_payload(payload, faculty, program)


def build_raw_sheets_from_api(
    spreadsheet_id: str,
    api_key: str,
    faculty: Optional[str] = None,
    program: Optional[str] = None,
    sheet_filter_regex: str = "курс",
) -> tuple[list[RawSheet], Optional[int]]:
    """Обратная совместимость: (raw_sheets, academic_year_start)."""
    r = fetch_schedule_from_api(spreadsheet_id, api_key, faculty, program, sheet_filter_regex)
    return r.raw_sheets, r.academic_year_start
